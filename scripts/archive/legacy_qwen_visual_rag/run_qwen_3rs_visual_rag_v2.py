import argparse
import json
import math
import re
import time
import traceback
from pathlib import Path

import cv2
import numpy as np
import pandas as pd
import torch
from PIL import Image, ImageDraw
from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from transformers import AutoProcessor, Qwen2_5_VLForConditionalGeneration
from qwen_vl_utils import process_vision_info


TARGETS = ["WD_P", "WD_T", "CF_P", "CF_T"]


# ----------------------------
# Manual loading and chunking
# ----------------------------

def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_into_chunks(text, max_chars=1200, overlap_chars=150):
    text = clean_text(text)
    if not text:
        return []

    parts = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(parts) <= 1:
        parts = [p.strip() for p in text.splitlines() if p.strip()]

    chunks = []
    buf = ""

    for part in parts:
        candidate = f"{buf}\n{part}".strip()

        if len(candidate) <= max_chars:
            buf = candidate
            continue

        if buf:
            chunks.append(buf)
            tail = buf[-overlap_chars:]
            buf = f"{tail}\n{part}".strip()
        else:
            chunks.append(part[:max_chars])
            buf = part[max_chars - overlap_chars:]

    if buf:
        chunks.append(buf)

    return chunks


def read_official_manual(pdf_path):
    reader = PdfReader(str(pdf_path))
    chunks = []

    for page_num, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""

        for chunk_text in split_into_chunks(page_text):
            chunks.append({
                "source": "official_manual",
                "page": page_num,
                "text": chunk_text,
            })

    return chunks


def read_one_minute_guide(docx_path):
    doc = Document(str(docx_path))
    blocks = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            blocks.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(x for x in cells if x)
            if line:
                blocks.append(line)

    text = "\n\n".join(blocks)

    return [
        {
            "source": "one_minute_guide",
            "page": None,
            "text": chunk_text,
        }
        for chunk_text in split_into_chunks(text)
    ]


def chunk_tags(text):
    t = text.lower()
    tags = set()

    if any(x in t for x in [
        "withdraw", "movement away", "moves away", "moving away",
        "shuts down", "minimal response", "avoidant", "masking",
        "content/affect", "deferential", "appeasing"
    ]):
        tags.add("withdrawal")

    if any(x in t for x in [
        "confront", "movement against", "moves against", "moving against",
        "criticiz", "complain", "push back", "control", "pressure",
        "hostil", "interrupt"
    ]):
        tags.add("confrontation")

    if any(x in t for x in [
        "working together", "alliance", "collaboration", "collaborating",
        "bond", "engaged in the work"
    ]):
        tags.add("alliance")

    if any(x in t for x in [
        "rating", "salient", "salience", "dominance",
        "1-minute", "1 minute", "five-point", "5-point"
    ]):
        tags.add("scoring")

    if any(x in t for x in [
        "not every", "not all", "do not assume", "do not confuse",
        "distinguish", "uncertain", "boundary", "caution",
        "should not", "does not constitute", "not a rupture"
    ]):
        tags.add("boundary")

    return sorted(tags)


class ManualRAG:
    def __init__(self, official_manual, one_minute_guide):
        chunks = read_official_manual(official_manual)
        chunks += read_one_minute_guide(one_minute_guide)

        for i, chunk in enumerate(chunks):
            chunk["chunk_id"] = i
            chunk["tags"] = chunk_tags(chunk["text"])

        self.chunks = chunks
        self.vectorizer = TfidfVectorizer(
            lowercase=True,
            stop_words="english",
            ngram_range=(1, 2),
            min_df=1,
            sublinear_tf=True,
        )
        self.matrix = self.vectorizer.fit_transform(
            [c["text"] for c in self.chunks]
        )

    def search(self, query, top_k=3, required_tag=None, source=None):
        q_vec = self.vectorizer.transform([query])
        cosine = cosine_similarity(q_vec, self.matrix).ravel()

        q_terms = set(re.findall(r"[a-zA-Z][a-zA-Z/_-]+", query.lower()))
        scored = []

        for idx, base_score in enumerate(cosine):
            chunk = self.chunks[idx]

            if required_tag and required_tag not in chunk["tags"]:
                continue

            if source and chunk["source"] != source:
                continue

            chunk_terms = set(
                re.findall(r"[a-zA-Z][a-zA-Z/_-]+", chunk["text"].lower())
            )

            lexical_overlap = (
                len(q_terms & chunk_terms) / max(1, len(q_terms))
            )

            tag_boost = 0.08 if required_tag in chunk["tags"] else 0.0
            official_boost = 0.03 if chunk["source"] == "official_manual" else 0.0

            final_score = (
                float(base_score)
                + 0.20 * lexical_overlap
                + tag_boost
                + official_boost
            )

            scored.append((final_score, chunk))

        scored.sort(key=lambda x: x[0], reverse=True)

        results = []
        for score, chunk in scored[:top_k]:
            item = dict(chunk)
            item["retrieval_score"] = round(score, 5)
            results.append(item)

        return results

    def retrieve(self, observation, top_k=2):
        obs = json.dumps(observation, ensure_ascii=False)

        target_queries = {
            "WD_P": (
                "patient movement away withdrawal shut down minimal response "
                "avoids masking content affect split deferential appeasing "
                "nonverbal gaze posture head turn orientation " + obs
            ),
            "WD_T": (
                "therapist movement away withdrawal shut down minimal response "
                "avoids masking content affect split deferential appeasing "
                "nonverbal gaze posture head turn orientation " + obs
            ),
            "CF_P": (
                "patient movement against confrontation criticism complaint "
                "pushing back pressure control interruption hostility "
                "power struggle autonomy " + obs
            ),
            "CF_T": (
                "therapist movement against confrontation criticism complaint "
                "pushing back pressure control interruption hostility "
                "power struggle autonomy " + obs
            ),
        }

        result = {}

        for target, query in target_queries.items():
            tag = "withdrawal" if target.startswith("WD") else "confrontation"
            result[target] = self.search(
                query,
                top_k=top_k,
                required_tag=tag,
            )

        boundary_query = (
            "counterexample boundary not every pause brief response smile laugh "
            "straight face gaze look away story disagreement interruption "
            "self assertion normal therapy does not constitute rupture"
        )

        result["boundaries"] = self.search(
            boundary_query,
            top_k=4,
            required_tag="boundary",
        )

        result["one_minute_scoring"] = self.search(
            "1-minute rupture coding rating 1 2 3 4 5 clear marker "
            "unclear recognizable defensible dominance salience",
            top_k=4,
            required_tag="scoring",
            source="one_minute_guide",
        )

        result["alliance_context"] = self.search(
            "working together collaboration alliance bond therapy work "
            "before identifying rupture gauge working alliance",
            top_k=2,
            required_tag="alliance",
        )

        return result


# ----------------------------
# Video sampling
# ----------------------------

def resize_keep_aspect(image, target_width):
    width, height = image.size

    if width <= target_width:
        return image

    scale = target_width / width
    target_height = max(1, int(round(height * scale)))

    return image.resize(
        (target_width, target_height),
        Image.Resampling.LANCZOS,
    )


def add_timestamp(image, timestamp):
    image = image.copy()
    draw = ImageDraw.Draw(image)

    label = f"{timestamp:05.1f}s"
    box = (8, 8, 92, 31)

    draw.rectangle(box, fill="black")
    draw.text((13, 12), label, fill="white")

    return image


def sample_video_frames(
    video_path,
    sample_fps=0.75,
    max_frames=48,
    frame_width=448,
):
    cap = cv2.VideoCapture(str(video_path))

    if not cap.isOpened():
        raise RuntimeError(f"Could not open video: {video_path}")

    fps = cap.get(cv2.CAP_PROP_FPS)
    frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))

    if not fps or fps <= 0:
        fps = 25.0

    duration = frame_count / fps if frame_count > 0 else 60.0

    wanted = max(2, int(round(duration * sample_fps)))
    n_frames = min(max_frames, wanted)

    timestamps = np.linspace(
        0,
        max(0.0, duration - 0.001),
        n_frames,
    )

    frames = []

    for ts in timestamps:
        cap.set(cv2.CAP_PROP_POS_MSEC, float(ts * 1000))
        ok, frame = cap.read()

        if not ok:
            continue

        frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        image = Image.fromarray(frame)
        image = resize_keep_aspect(image, frame_width)
        image = add_timestamp(image, float(ts))

        frames.append(image)

    cap.release()

    if not frames:
        raise RuntimeError(f"No frames extracted from {video_path}")

    return frames


# ----------------------------
# Qwen
# ----------------------------

def extract_json(text):
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)

    start = text.find("{")
    end = text.rfind("}")

    if start < 0 or end < 0:
        raise ValueError(f"No JSON object found:\n{text}")

    return json.loads(text[start:end + 1])


class Qwen3RS:
    def __init__(
        self,
        model_id,
        observer_fps,
        max_frames,
        frame_width,
    ):
        self.model_id = model_id
        self.observer_fps = observer_fps
        self.max_frames = max_frames
        self.frame_width = frame_width

        if torch.cuda.is_available():
            dtype = (
                torch.bfloat16
                if torch.cuda.is_bf16_supported()
                else torch.float16
            )
        else:
            dtype = torch.float32

        print("CUDA available:", torch.cuda.is_available(), flush=True)
        print("Model dtype:", dtype, flush=True)

        self.processor = AutoProcessor.from_pretrained(model_id)

        self.model = Qwen2_5_VLForConditionalGeneration.from_pretrained(
            model_id,
            torch_dtype=dtype,
            device_map="auto",
            attn_implementation="sdpa",
        )

        self.model.eval()

        print("Model loaded", flush=True)
        print(
            "Model device:",
            next(self.model.parameters()).device,
            flush=True,
        )

    @property
    def device(self):
        return next(self.model.parameters()).device

    def generate(self, inputs, max_new_tokens):
        inputs = inputs.to(self.device)

        with torch.inference_mode():
            generated = self.model.generate(
                **inputs,
                max_new_tokens=max_new_tokens,
                do_sample=False,
                use_cache=True,
            )

        trimmed = [
            output_ids[len(input_ids):]
            for input_ids, output_ids
            in zip(inputs.input_ids, generated)
        ]

        return self.processor.batch_decode(
            trimmed,
            skip_special_tokens=True,
            clean_up_tokenization_spaces=False,
        )[0]

    def observe(self, video_path, role_description):
        frames = sample_video_frames(
            video_path,
            sample_fps=self.observer_fps,
            max_frames=self.max_frames,
            frame_width=self.frame_width,
        )

        prompt = f"""
You are a neutral VISUAL OBSERVER of a psychotherapy video.

ROLE LAYOUT:
{role_description}

This is observation stage only.

Describe only what can be seen in the sampled frames.
Do NOT decide whether anything is a rupture.
Do NOT label behavior as withdrawal, confrontation, repair, avoidance,
defensiveness, hostility, engagement, disengagement, anxiety, sadness,
agreement, disagreement, or any other psychological interpretation.

Do not infer speech content.
Do not write "before answering", "after being asked", or similar claims
unless the relationship is visually evident from the turn sequence.

You may describe:
- head turns
- gaze direction and changes
- posture and posture changes
- body orientation
- nodding and head shaking
- visible smiling or laughing
- hand and arm gestures
- visible mouth movement
- stillness
- entering or leaving the frame
- movement toward or away from the other person
- whether one person visually interrupts another person's ongoing gesture
- changes in visible interaction across the minute

Use the visible timestamps.

Return JSON only:

{{
  "visibility": {{
    "patient": "good|partial|poor|not_visible",
    "therapist": "good|partial|poor|not_visible"
  }},
  "events": [
    {{
      "start_sec": 0.0,
      "end_sec": 1.0,
      "actor": "patient|therapist|both|unknown",
      "behavior": "literal visual description",
      "interaction_context": "literal visual context only",
      "certainty": 0.0
    }}
  ],
  "patient_pattern": "short literal summary",
  "therapist_pattern": "short literal summary",
  "dyadic_visual_pattern": "short literal summary",
  "minute_summary": "short literal visual summary"
}}
""".strip()

        messages = [{
            "role": "user",
            "content": [
                {
                    "type": "video",
                    "video": frames,
                    "fps": float(self.observer_fps),
                },
                {
                    "type": "text",
                    "text": prompt,
                },
            ],
        }]

        chat_text = self.processor.apply_chat_template(
            messages,
            tokenize=False,
            add_generation_prompt=True,
        )

        image_inputs, video_inputs, video_kwargs = process_vision_info(
            messages,
            return_video_kwargs=True,
        )

        video_kwargs["fps"] = float(self.observer_fps)
        video_kwargs["do_sample_frames"] = False

        inputs = self.processor(
            text=[chat_text],
            images=image_inputs,
            videos=video_inputs,
            padding=True,
            return_tensors="pt",
            **video_kwargs,
        )

        print(
            f"Observer frames: {len(frames)} "
            f"at width <= {self.frame_width}px",
            flush=True,
        )

        raw = self.generate(inputs, max_new_tokens=900)
        return extract_json(raw), raw

    def judge(
        self,
        observation,
        retrieved,
        role_description,
    ):
        evidence_text = format_retrieved_evidence(retrieved)

        prompt = f"""
You are a 3RS v2022 coding judge for one 1-minute psychotherapy segment.

ROLE LAYOUT:
{role_description}

You are NOT watching the video in this stage.

You receive:
1. a neutral visual observation generated independently of 3RS coding
2. passages retrieved from the official 3RS v2022 manual
3. passages retrieved from the project 1-minute coding guide

SOURCE PRIORITY:
- Official 3RS v2022 manual is authoritative.
- The 1-minute guide supplies project-specific operational rules.
- If they conflict, follow the official manual.

IMPORTANT LIMITATIONS:
- This experiment is visual-only.
- Do not invent speech, wording, tone of voice, or meaning.
- A visually similar behavior is NOT automatically a rupture marker.
- Use retrieved boundary/counterexamples to avoid overcoding.
- Score WD_P, WD_T, CF_P, and CF_T independently.
- Confidence means confidence in the rating, not rupture intensity.

NEUTRAL VISUAL OBSERVATION:
{json.dumps(observation, indent=2, ensure_ascii=False)}

RETRIEVED 3RS MATERIAL:
{evidence_text}

For each target:
A. identify candidate manual marker(s)
B. identify supporting visual event(s)
C. identify counterevidence / missing context
D. decide whether the marker is sufficiently supported
E. apply the 1-minute scoring rules

Return JSON only:

{{
  "scores": {{
    "WD_P": 1,
    "WD_T": 1,
    "CF_P": 1,
    "CF_T": 1
  }},
  "candidate_markers": {{
    "WD_P": [],
    "WD_T": [],
    "CF_P": [],
    "CF_T": []
  }},
  "supporting_visual_events": {{
    "WD_P": [],
    "WD_T": [],
    "CF_P": [],
    "CF_T": []
  }},
  "counterevidence": {{
    "WD_P": [],
    "WD_T": [],
    "CF_P": [],
    "CF_T": []
  }},
  "matched_manual_chunk_ids": {{
    "WD_P": [],
    "WD_T": [],
    "CF_P": [],
    "CF_T": []
  }},
  "confidence": {{
    "WD_P": 0.0,
    "WD_T": 0.0,
    "CF_P": 0.0,
    "CF_T": 0.0
  }},
  "information_limitations": [],
  "primary_class": "none|WD_P|WD_T|CF_P|CF_T|mixed",
  "reasoning_summary": "short evidence-based explanation"
}}

SCORING PRINCIPLE:
- 1 = no identifiable supported rupture marker
- 2 = recognizable rupture direction but unclear / incomplete marker
- 3 = clear defensible rupture marker
- 4 or 5 = only when the retrieved 1-minute rules justify
  greater salience/dominance
""".strip()

        messages = [{
            "role": "user",
            "content": [{
                "type": "text",
                "text": prompt,
            }],
        }]

        chat_text = self.processor.apply_chat_template(
            messages,
            tokenize=False,
            add_generation_prompt=True,
        )

        inputs = self.processor(
            text=[chat_text],
            padding=True,
            return_tensors="pt",
        )

        raw = self.generate(inputs, max_new_tokens=1100)
        result = extract_json(raw)

        for target in TARGETS:
            score = float(result["scores"][target])
            score = int(round(score))
            result["scores"][target] = min(5, max(1, score))

            conf = float(
                result.get("confidence", {}).get(target, 0.0)
            )
            result["confidence"][target] = min(1.0, max(0.0, conf))

        return result, raw


# ----------------------------
# Formatting and output
# ----------------------------

def format_retrieved_evidence(retrieved):
    lines = []

    for group in TARGETS + [
        "boundaries",
        "one_minute_scoring",
        "alliance_context",
    ]:
        lines.append(f"\n### {group}")

        for item in retrieved.get(group, []):
            page = item.get("page")
            location = (
                f"page {page}"
                if page is not None
                else "project guide"
            )

            lines.append(
                f"[chunk_id={item['chunk_id']} | "
                f"source={item['source']} | "
                f"{location} | "
                f"retrieval_score={item['retrieval_score']}]\n"
                f"{item['text']}"
            )

    return "\n".join(lines)


def compact_retrieval(retrieved):
    output = {}

    for group, items in retrieved.items():
        output[group] = []

        for item in items:
            output[group].append({
                "chunk_id": item["chunk_id"],
                "source": item["source"],
                "page": item["page"],
                "tags": item["tags"],
                "retrieval_score": item["retrieval_score"],
                "text": item["text"],
            })

    return output


def flatten_result(
    segment_idx,
    segment_path,
    observation,
    judgment,
    elapsed_sec,
):
    row = {
        "segment_idx": int(segment_idx),
        "segment_path": str(segment_path),
        "status": "ok",
        "error": "",
        "primary_class": judgment.get("primary_class", ""),
        "observation_summary": observation.get("minute_summary", ""),
        "reasoning_summary": judgment.get("reasoning_summary", ""),
        "information_limitations": json.dumps(
            judgment.get("information_limitations", []),
            ensure_ascii=False,
        ),
        "elapsed_sec": round(float(elapsed_sec), 3),
    }

    for target in TARGETS:
        row[f"{target}_pred"] = judgment["scores"].get(target)
        row[f"{target}_conf"] = judgment.get(
            "confidence",
            {},
        ).get(target)

        row[f"{target}_markers"] = json.dumps(
            judgment.get(
                "candidate_markers",
                {},
            ).get(target, []),
            ensure_ascii=False,
        )

        row[f"{target}_manual_chunks"] = json.dumps(
            judgment.get(
                "matched_manual_chunk_ids",
                {},
            ).get(target, []),
            ensure_ascii=False,
        )

    return row


def save_csv(rows, path):
    if not rows:
        return

    df = pd.DataFrame(rows)

    if "segment_idx" in df.columns:
        df["segment_idx"] = pd.to_numeric(
            df["segment_idx"],
            errors="coerce",
        )

        df = (
            df.dropna(subset=["segment_idx"])
            .sort_values("segment_idx")
            .drop_duplicates("segment_idx", keep="last")
            .reset_index(drop=True)
        )

    df.to_csv(
        path,
        index=False,
        encoding="utf-8-sig",
    )


def parse_segment_indices(value):
    if not value:
        return None

    return {
        int(x.strip())
        for x in value.split(",")
        if x.strip()
    }


def run(args):
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    csv_path = output_dir / "qwen_3rs_visual_rag_predictions.csv"
    jsonl_path = output_dir / "qwen_3rs_visual_rag_details.jsonl"
    chunks_path = output_dir / "manual_rag_chunks.json"

    print("Building RAG index...", flush=True)

    rag = ManualRAG(
        args.official_manual,
        args.one_minute_guide,
    )

    chunks_path.write_text(
        json.dumps(
            rag.chunks,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print(
        f"Indexed {len(rag.chunks)} manual chunks",
        flush=True,
    )

    qwen = Qwen3RS(
        model_id=args.model_id,
        observer_fps=args.observer_fps,
        max_frames=args.max_frames,
        frame_width=args.frame_width,
    )

    segments = pd.read_csv(args.segments_csv)

    required = {"segment_idx", "segment_path"}
    missing = required - set(segments.columns)

    if missing:
        raise ValueError(
            f"segments CSV is missing columns: {sorted(missing)}"
        )

    segments["segment_idx"] = pd.to_numeric(
        segments["segment_idx"],
        errors="raise",
    ).astype(int)

    requested_indices = parse_segment_indices(
        args.segment_indices
    )

    if requested_indices is not None:
        segments = segments[
            segments["segment_idx"].isin(requested_indices)
        ].copy()

    if args.max_segments is not None:
        segments = segments.head(args.max_segments)

    rows = []
    completed = set()

    if csv_path.exists():
        old = pd.read_csv(csv_path)

        if not old.empty:
            rows = old.to_dict("records")

            if "status" in old.columns:
                completed = set(
                    pd.to_numeric(
                        old.loc[
                            old["status"] == "ok",
                            "segment_idx",
                        ],
                        errors="coerce",
                    )
                    .dropna()
                    .astype(int)
                )

    total = len(segments)

    print(f"Segments in run: {total}", flush=True)
    print(f"CSV: {csv_path}", flush=True)
    print(f"JSONL: {jsonl_path}", flush=True)

    for pos, row in enumerate(
        segments.itertuples(index=False),
        start=1,
    ):
        segment_idx = int(row.segment_idx)
        segment_path = Path(row.segment_path)

        if segment_idx in completed:
            print(
                f"[{pos}/{total}] segment {segment_idx}: already done",
                flush=True,
            )
            continue

        print(
            f"\n[{pos}/{total}] segment {segment_idx}",
            flush=True,
        )
        print(
            f"Processing: {segment_path.name}",
            flush=True,
        )

        started = time.time()

        try:
            print("Stage 1: visual observation...", flush=True)

            observation, observer_raw = qwen.observe(
                segment_path,
                args.role_description,
            )

            print(
                "Observation:",
                observation.get("minute_summary", ""),
                flush=True,
            )

            print("Stage 2: manual retrieval...", flush=True)

            retrieved = rag.retrieve(
                observation,
                top_k=args.top_k,
            )

            print("Stage 3: 3RS judgment...", flush=True)

            judgment, judge_raw = qwen.judge(
                observation,
                retrieved,
                args.role_description,
            )

            elapsed = time.time() - started

            result_row = flatten_result(
                segment_idx,
                segment_path,
                observation,
                judgment,
                elapsed,
            )

            rows = [
                r for r in rows
                if int(float(r["segment_idx"])) != segment_idx
            ]
            rows.append(result_row)

            save_csv(rows, csv_path)

            detail = {
                "segment_idx": segment_idx,
                "segment_path": str(segment_path),
                "observation": observation,
                "retrieval": compact_retrieval(retrieved),
                "judgment": judgment,
                "observer_raw": observer_raw,
                "judge_raw": judge_raw,
                "elapsed_sec": elapsed,
            }

            with jsonl_path.open(
                "a",
                encoding="utf-8",
            ) as f:
                f.write(
                    json.dumps(
                        detail,
                        ensure_ascii=False,
                    )
                    + "\n"
                )

            completed.add(segment_idx)

            print(
                "Scores:",
                judgment["scores"],
                flush=True,
            )
            print(
                "Primary:",
                judgment.get("primary_class"),
                flush=True,
            )
            print(
                f"Finished in {elapsed:.1f}s",
                flush=True,
            )

            if torch.cuda.is_available():
                torch.cuda.empty_cache()

        except Exception as exc:
            elapsed = time.time() - started

            print(
                f"ERROR segment {segment_idx}: {exc}",
                flush=True,
            )
            traceback.print_exc()

            error_row = {
                "segment_idx": segment_idx,
                "segment_path": str(segment_path),
                "status": "error",
                "error": repr(exc),
                "elapsed_sec": round(elapsed, 3),
            }

            rows = [
                r for r in rows
                if int(float(r["segment_idx"])) != segment_idx
            ]
            rows.append(error_row)

            save_csv(rows, csv_path)

            if torch.cuda.is_available():
                torch.cuda.empty_cache()

    print("\nFinished", flush=True)

    ok = sum(
        1 for r in rows
        if r.get("status") == "ok"
    )
    errors = sum(
        1 for r in rows
        if r.get("status") == "error"
    )

    print(f"Successful predictions: {ok}", flush=True)
    print(f"Errors: {errors}", flush=True)
    print(f"Saved to: {csv_path}", flush=True)


def build_parser():
    parser = argparse.ArgumentParser(
        description=(
            "Visual observer -> 3RS manual RAG -> "
            "3RS judge pipeline"
        )
    )

    parser.add_argument(
        "--segments-csv",
        required=True,
    )
    parser.add_argument(
        "--official-manual",
        required=True,
    )
    parser.add_argument(
        "--one-minute-guide",
        required=True,
    )
    parser.add_argument(
        "--output-dir",
        default="./qwen_3rs_visual_rag_results",
    )

    parser.add_argument(
        "--model-id",
        default="Qwen/Qwen2.5-VL-3B-Instruct",
    )

    parser.add_argument(
        "--observer-fps",
        type=float,
        default=0.75,
    )
    parser.add_argument(
        "--max-frames",
        type=int,
        default=48,
    )
    parser.add_argument(
        "--frame-width",
        type=int,
        default=448,
    )

    parser.add_argument(
        "--top-k",
        type=int,
        default=2,
    )

    parser.add_argument(
        "--max-segments",
        type=int,
        default=None,
    )

    parser.add_argument(
        "--segment-indices",
        default=None,
        help="Comma-separated zero-based segment indices, e.g. 0,5,28",
    )

    parser.add_argument(
        "--role-description",
        default=(
            "Patient is the person on the LEFT side of the video. "
            "Therapist is the person on the RIGHT side of the video."
        ),
    )

    return parser


if __name__ == "__main__":
    args = build_parser().parse_args()
    run(args)