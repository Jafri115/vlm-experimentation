import argparse
import json
import math
import re
import sys
import traceback
from pathlib import Path

import cv2
import numpy as np
import pandas as pd
import torch
from docx import Document
from pypdf import PdfReader
from PIL import Image

from transformers import (
    AutoProcessor,
    Qwen2_5_VLForConditionalGeneration,
)
from qwen_vl_utils import process_vision_info


# ============================================================
# Configuration
# ============================================================

MODEL_ID = "Qwen/Qwen2.5-VL-3B-Instruct"

# One frame every 2 seconds -> about 30 frames for a 60-second segment
VLM_VIDEO_FPS = 0.5
MAX_VIDEO_FRAMES = 32

MAX_NEW_TOKENS = 550


# ============================================================
# Load Qwen
# ============================================================

if torch.cuda.is_available():
    if torch.cuda.is_bf16_supported():
        model_dtype = torch.bfloat16
    else:
        model_dtype = torch.float16
else:
    model_dtype = torch.float32

print("CUDA available:", torch.cuda.is_available())
print("Model dtype:", model_dtype)


processor = AutoProcessor.from_pretrained(
    MODEL_ID
)


model = Qwen2_5_VLForConditionalGeneration.from_pretrained(
    MODEL_ID,
    torch_dtype=model_dtype,
    device_map="auto",
    attn_implementation="sdpa",
)

model.eval()

model_device = next(model.parameters()).device

print("Model loaded")
print("Model device:", model_device)


# ============================================================
# Sample frames directly with OpenCV
#
# This bypasses:
# - TorchCodec
# - torchvision.io.read_video
# ============================================================

def sample_video_frames(
    video_path: Path,
    sample_fps: float = 0.5,
    max_frames: int = 32,
):
    video_path = Path(video_path)

    cap = cv2.VideoCapture(str(video_path))

    if not cap.isOpened():
        raise RuntimeError(
            f"Could not open video: {video_path}"
        )

    source_fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(
        cap.get(cv2.CAP_PROP_FRAME_COUNT)
    )

    if source_fps is None or source_fps <= 0:
        source_fps = 25.0

    duration = total_frames / source_fps

    requested_frames = max(
        2,
        int(round(duration * sample_fps))
    )

    n_frames = min(
        max_frames,
        requested_frames
    )

    timestamps = np.linspace(
        0,
        max(0, duration - 0.001),
        n_frames,
    )

    frames = []

    for timestamp in timestamps:

        cap.set(
            cv2.CAP_PROP_POS_MSEC,
            float(timestamp * 1000),
        )

        ok, frame = cap.read()

        if not ok:
            continue

        frame = cv2.cvtColor(
            frame,
            cv2.COLOR_BGR2RGB,
        )

        pil_frame = Image.fromarray(frame)

        frames.append(pil_frame)

    cap.release()

    if len(frames) == 0:
        raise RuntimeError(
            f"No frames could be extracted from: {video_path}"
        )

    return frames


# ============================================================
# Helper for feature JSON
# ============================================================

def compact_features(features: dict) -> dict:

    keep = [
        "patient_motion_mean",
        "therapist_motion_mean",
        "patient_face_visible_fraction",
        "therapist_face_visible_fraction",
        "patient_face_area_fraction_mean",
        "therapist_face_area_fraction_mean",
        "patient_face_center_jitter",
        "therapist_face_center_jitter",
        "brightness_mean_0_255",
        "blur_laplacian_var_mean",
    ]

    result = {}

    for key in keep:

        value = features.get(key)

        if isinstance(
            value,
            (float, np.floating)
        ):

            if np.isnan(value):
                result[key] = None
            else:
                result[key] = round(
                    float(value),
                    4
                )

        elif isinstance(
            value,
            (int, np.integer)
        ):
            result[key] = int(value)

        else:
            result[key] = value

    return result


# ============================================================
# Full 3RS manual loading
# ============================================================

ROLE_DESCRIPTION = """
Patient is the person on the LEFT side of the video.
Therapist is the person on the RIGHT side of the video.
"""


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(
            f"\n--- OFFICIAL MANUAL PAGE {page_num} ---\n{text.strip()}"
        )
    return "\n".join(pages)


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_num, table in enumerate(doc.tables, start=1):
        parts.append(f"\n--- PROJECT GUIDE TABLE {table_num} ---")
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def load_3rs_reference(official_manual: Path, project_guide: Path) -> str:
    if not official_manual.exists():
        raise FileNotFoundError(f"Official manual not found: {official_manual}")
    if not project_guide.exists():
        raise FileNotFoundError(f"Project guide not found: {project_guide}")

    official_text = extract_pdf_text(official_manual)
    guide_text = extract_docx_text(project_guide)

    reference = f"""
3RS REFERENCE PACKAGE

SOURCE PRIORITY
1. The official 3RS v2022 manual by Eubanks & Muran is the governing source.
2. The project 1-minute coding guide operationalizes the official manual for 1-minute segments.
3. If the project guide conflicts with the official manual, follow the official manual.

============================================================
OFFICIAL 3RS v2022 MANUAL - COMPLETE EXTRACTED TEXT
============================================================

{official_text}

============================================================
PROJECT 1-MINUTE CODING GUIDE - COMPLETE EXTRACTED TEXT
============================================================

{guide_text}
"""

    return reference.strip()


# ============================================================
# Build prompt from the full manuals
# ============================================================

def build_prompt(features: dict, reference_text: str) -> str:
    feature_json = json.dumps(
        compact_features(features),
        indent=2,
    )

    prompt = f"""
You are acting as a 3RS coder for one 60-second psychotherapy segment.

Read and apply the COMPLETE reference material below. Do not substitute a
shortened definition of rupture for the manual. Base the coding decision on
the official 3RS v2022 manual and the project-specific 1-minute guide.

IMPORTANT CODING INSTRUCTIONS

- First use the official manual to understand rupture, withdrawal,
  confrontation, salience, markers, subtypes, boundaries, and examples.
- Then use the project guide for the 1-minute operational rules.
- The official manual wins if the two sources conflict.
- Do not default to score 1 simply because the interaction looks generally
  calm, cooperative, or engaged. The manual allows rupture markers to occur
  within otherwise collaborative interaction.
- For every score above 1, identify the manual marker or subtype that supports
  it before deciding the final score.
- For every score of 1, actively check whether any manual-described marker is
  present before concluding that no meaningful marker is visible.
- Rate Patient Withdrawal (WD_P), Therapist Withdrawal (WD_T), Patient
  Confrontation (CF_P), and Therapist Confrontation (CF_T) separately.

MODALITY LIMITATION

You are receiving sampled VIDEO FRAMES only. You cannot hear the audio and you
have no transcript in this experiment. The manual contains verbal as well as
nonverbal criteria. Apply verbal criteria only when they can genuinely be
supported by visible behavior. Do not invent spoken content, topic changes,
minimal verbal responses, tone, disagreement, or other audio/text information.
Use nonverbal and interactional evidence that is actually visible. If a marker
cannot be judged because required verbal information is unavailable, state this
as ambiguity rather than assuming it occurred.

ROLE LAYOUT

{ROLE_DESCRIPTION}

AUXILIARY VIDEO FEATURES

{feature_json}

The numeric features are weak supporting measurements only. They are not 3RS
labels and must not override the manual or the visible interaction.

============================================================
COMPLETE 3RS REFERENCE MATERIAL
============================================================

{reference_text}

============================================================
CODING TASK FOR THIS SEGMENT
============================================================

After reading the reference material and observing the video frames:

1. Check the segment for the specific withdrawal markers described in the
   manual for the patient and therapist.
2. Check the segment for the specific confrontation markers described in the
   manual for the patient and therapist.
3. Distinguish genuine rupture markers from ordinary therapy behavior using
   the manual's boundary rules.
4. Apply the 1-minute guide's scoring rules and salience anchors.
5. Return the four final 1-5 ratings.

Return ONLY valid JSON using exactly this structure:

{{
    "scores": {{
        "WD_P": 1,
        "WD_T": 1,
        "CF_P": 1,
        "CF_T": 1
    }},
    "primary_class": "none",
    "confidence": {{
        "WD_P": 0.0,
        "WD_T": 0.0,
        "CF_P": 0.0,
        "CF_T": 0.0
    }},
    "manual_marker": {{
        "WD_P": [],
        "WD_T": [],
        "CF_P": [],
        "CF_T": []
    }},
    "visible_evidence": {{
        "WD_P": [],
        "WD_T": [],
        "CF_P": [],
        "CF_T": []
    }},
    "counterevidence_or_ambiguity": [],
    "visibility_notes": "",
    "segment_summary": ""
}}

Rules for the JSON:

- Scores must be integers from 1 to 5.
- Confidence must be between 0.0 and 1.0.
- manual_marker should name the relevant marker/subtype from the manual when
  one is observed; otherwise use an empty list.
- visible_evidence must describe only behavior actually visible in the frames.
- primary_class must be one of: none, WD_P, WD_T, CF_P, CF_T, mixed.
- Use mixed when multiple rupture dimensions are meaningfully salient.
- Use none only after checking the manual markers and finding no meaningful
  evidence above the score-1 threshold.
"""

    return prompt.strip()


# ============================================================
# Parse Qwen JSON
# ============================================================

def extract_json_object(text: str) -> dict:

    text = text.strip()

    if text.startswith("```"):

        text = re.sub(
            r"^```(?:json)?\s*",
            "",
            text,
            flags=re.IGNORECASE,
        )

        text = re.sub(
            r"\s*```$",
            "",
            text,
        )

    start = text.find("{")
    end = text.rfind("}")

    if start == -1 or end == -1:
        raise ValueError(
            "Could not find JSON in model response:\n"
            + text
        )

    return json.loads(
        text[start:end + 1]
    )


# ============================================================
# Validate model output
# ============================================================

def validate_prediction(prediction: dict) -> dict:

    targets = [
        "WD_P",
        "WD_T",
        "CF_P",
        "CF_T",
    ]

    scores = prediction.get(
        "scores",
        {}
    )

    confidence = prediction.get(
        "confidence",
        {}
    )

    for target in targets:

        if target not in scores:
            raise ValueError(
                f"Missing score: {target}"
            )

        score = int(
            round(
                float(scores[target])
            )
        )

        scores[target] = min(
            5,
            max(1, score)
        )

        conf = float(
            confidence.get(
                target,
                0.0
            )
        )

        confidence[target] = min(
            1.0,
            max(0.0, conf)
        )

    allowed_classes = {
        "none",
        "WD_P",
        "WD_T",
        "CF_P",
        "CF_T",
        "mixed",
    }

    primary_class = prediction.get(
        "primary_class",
        "none",
    )

    if primary_class not in allowed_classes:

        if max(scores.values()) >= 3:
            primary_class = "mixed"
        else:
            primary_class = "none"

    prediction["scores"] = scores
    prediction["confidence"] = confidence
    prediction["primary_class"] = primary_class

    return prediction


# ============================================================
# Main Qwen prediction function
# ============================================================

def predict_segment(
    segment_path: Path,
    features: dict,
    reference_text: str,
) -> tuple[dict, str]:

    segment_path = Path(segment_path)

    prompt = build_prompt(features, reference_text)

    frames = sample_video_frames(
        segment_path,
        sample_fps=VLM_VIDEO_FPS,
        max_frames=MAX_VIDEO_FRAMES,
    )

    print(
        f"Processing: {segment_path.name}"
    )

    print(
        f"Frames sent to Qwen: {len(frames)}"
    )

    print(
        f"Temporal sampling FPS: {VLM_VIDEO_FPS}"
    )

    messages = [
        {
            "role": "user",
            "content": [
                {
                    "type": "video",

                    # Important:
                    # list of PIL frames avoids
                    # TorchCodec / torchvision reader
                    "video": frames,

                    "fps": float(
                        VLM_VIDEO_FPS
                    ),
                },

                {
                    "type": "text",
                    "text": prompt,
                },
            ],
        }
    ]

    chat_text = processor.apply_chat_template(
        messages,
        tokenize=False,
        add_generation_prompt=True,
    )

    image_inputs, video_inputs, video_kwargs = (
        process_vision_info(
            messages,
            return_video_kwargs=True,
        )
    )

    # Newer Transformers expects scalar fps.
    # qwen-vl-utils may return [0.5].
    video_kwargs["fps"] = float(
        VLM_VIDEO_FPS
    )

    # Frames are already sampled manually.
    video_kwargs["do_sample_frames"] = False

    print(
        "Video processor kwargs:",
        video_kwargs,
    )

    inputs = processor(
        text=[chat_text],

        images=image_inputs,

        videos=video_inputs,

        padding=True,

        return_tensors="pt",

        # Frames have already been prepared.
        do_resize=False,

        **video_kwargs,
    )

    # Avoid dependency on a global model_device variable.
    device = next(
        model.parameters()
    ).device

    inputs = inputs.to(device)

    print(
        "Running Qwen inference..."
    )

    with torch.inference_mode():

        generated_ids = model.generate(
            **inputs,
            max_new_tokens=MAX_NEW_TOKENS,
            do_sample=False,
        )

    trimmed_ids = []

    for input_ids, output_ids in zip(
        inputs.input_ids,
        generated_ids,
    ):

        trimmed_ids.append(
            output_ids[
                len(input_ids):
            ]
        )

    raw_text = processor.batch_decode(
        trimmed_ids,
        skip_special_tokens=True,
        clean_up_tokenization_spaces=False,
    )[0]

    print("\nRaw Qwen response:")
    print(raw_text)

    parsed = extract_json_object(
        raw_text
    )

    parsed = validate_prediction(
        parsed
    )

    return parsed, raw_text

# ============================================================
# Standalone batch runner
# ============================================================

TARGETS = ["WD_P", "WD_T", "CF_P", "CF_T"]


def flatten_prediction(segment_idx, segment_path, prediction, raw_text):
    scores = prediction.get("scores", {})
    confidence = prediction.get("confidence", {})
    evidence = prediction.get("visible_evidence", {})
    manual_marker = prediction.get("manual_marker", {})

    row = {
        "segment_idx": int(segment_idx),
        "segment_path": str(segment_path),
        "WD_P_pred": scores.get("WD_P"),
        "WD_T_pred": scores.get("WD_T"),
        "CF_P_pred": scores.get("CF_P"),
        "CF_T_pred": scores.get("CF_T"),
        "WD_P_conf": confidence.get("WD_P"),
        "WD_T_conf": confidence.get("WD_T"),
        "CF_P_conf": confidence.get("CF_P"),
        "CF_T_conf": confidence.get("CF_T"),
        "primary_class": prediction.get("primary_class"),
        "WD_P_manual_marker": json.dumps(manual_marker.get("WD_P", []), ensure_ascii=False),
        "WD_T_manual_marker": json.dumps(manual_marker.get("WD_T", []), ensure_ascii=False),
        "CF_P_manual_marker": json.dumps(manual_marker.get("CF_P", []), ensure_ascii=False),
        "CF_T_manual_marker": json.dumps(manual_marker.get("CF_T", []), ensure_ascii=False),
        "WD_P_evidence": json.dumps(evidence.get("WD_P", []), ensure_ascii=False),
        "WD_T_evidence": json.dumps(evidence.get("WD_T", []), ensure_ascii=False),
        "CF_P_evidence": json.dumps(evidence.get("CF_P", []), ensure_ascii=False),
        "CF_T_evidence": json.dumps(evidence.get("CF_T", []), ensure_ascii=False),
        "ambiguity": json.dumps(
            prediction.get("counterevidence_or_ambiguity", []),
            ensure_ascii=False,
        ),
        "visibility_notes": prediction.get("visibility_notes", ""),
        "segment_summary": prediction.get("segment_summary", ""),
        "raw_response": raw_text,
        "status": "ok",
        "error": "",
    }
    return row


def save_checkpoint(rows, csv_path):
    if not rows:
        return

    df = pd.DataFrame(rows)
    df = (
        df.sort_values("segment_idx")
        .drop_duplicates("segment_idx", keep="last")
        .reset_index(drop=True)
    )
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")


def run_batch(segments_csv, features_csv, output_dir, official_manual, project_guide, max_segments=None):
    segments_csv = Path(segments_csv)
    features_csv = Path(features_csv)
    output_dir = Path(output_dir)
    official_manual = Path(official_manual)
    project_guide = Path(project_guide)

    print("Loading complete 3RS reference material...")
    reference_text = load_3rs_reference(official_manual, project_guide)
    print(f"3RS reference loaded: {len(reference_text):,} characters")
    output_dir.mkdir(parents=True, exist_ok=True)

    predictions_csv = output_dir / "qwen_3rs_predictions.csv"
    predictions_jsonl = output_dir / "qwen_3rs_predictions.jsonl"

    segments_df = pd.read_csv(segments_csv)
    features_df = pd.read_csv(features_csv)

    required_segments = {"segment_idx", "segment_path"}
    missing = required_segments - set(segments_df.columns)
    if missing:
        raise ValueError(
            f"Segments CSV is missing required columns: {sorted(missing)}"
        )

    if "segment_idx" not in features_df.columns:
        raise ValueError("Features CSV must contain a 'segment_idx' column.")

    segments_df["segment_idx"] = pd.to_numeric(
        segments_df["segment_idx"], errors="raise"
    ).astype(int)
    features_df["segment_idx"] = pd.to_numeric(
        features_df["segment_idx"], errors="raise"
    ).astype(int)

    feature_lookup = (
        features_df.drop_duplicates("segment_idx", keep="last")
        .set_index("segment_idx")
    )

    rows = []
    completed = set()

    if predictions_csv.exists():
        previous = pd.read_csv(predictions_csv)
        if not previous.empty and "segment_idx" in previous.columns:
            previous["segment_idx"] = pd.to_numeric(
                previous["segment_idx"], errors="coerce"
            )
            previous = previous.dropna(subset=["segment_idx"])
            previous["segment_idx"] = previous["segment_idx"].astype(int)
            rows = previous.to_dict("records")

            # Retry previous errors, skip only successful predictions.
            if "status" in previous.columns:
                completed = set(
                    previous.loc[
                        previous["status"].fillna("") == "ok",
                        "segment_idx",
                    ].tolist()
                )
            else:
                completed = set(previous["segment_idx"].tolist())

            print(
                f"Loaded checkpoint: {len(previous)} rows, "
                f"{len(completed)} completed segments."
            )

    work = segments_df.sort_values("segment_idx").copy()
    if max_segments is not None:
        work = work.head(max_segments)

    total = len(work)
    print(f"Segments in run: {total}")
    print(f"Predictions CSV: {predictions_csv}")
    print(f"Predictions JSONL: {predictions_jsonl}")

    for position, segment in enumerate(work.itertuples(index=False), start=1):
        segment_idx = int(segment.segment_idx)

        if segment_idx in completed:
            print(f"[{position}/{total}] segment {segment_idx}: already done")
            continue

        if segment_idx not in feature_lookup.index:
            print(f"[{position}/{total}] segment {segment_idx}: missing features")
            error_row = {
                "segment_idx": segment_idx,
                "segment_path": str(segment.segment_path),
                "status": "error",
                "error": "Missing features for segment",
            }
            rows = [r for r in rows if int(r["segment_idx"]) != segment_idx]
            rows.append(error_row)
            save_checkpoint(rows, predictions_csv)
            continue

        features = feature_lookup.loc[segment_idx].to_dict()

        try:
            print(f"\n[{position}/{total}] segment {segment_idx}")

            prediction, raw_text = predict_segment(
                Path(segment.segment_path),
                features,
                reference_text,
            )

            result_row = flatten_prediction(
                segment_idx,
                segment.segment_path,
                prediction,
                raw_text,
            )

            rows = [r for r in rows if int(r["segment_idx"]) != segment_idx]
            rows.append(result_row)
            save_checkpoint(rows, predictions_csv)

            with predictions_jsonl.open("a", encoding="utf-8") as f:
                record = {
                    "segment_idx": segment_idx,
                    "segment_path": str(segment.segment_path),
                    "prediction": prediction,
                    "raw_response": raw_text,
                }
                f.write(json.dumps(record, ensure_ascii=False) + "\n")

            completed.add(segment_idx)
            print(f"Saved segment {segment_idx}")

        except Exception as e:
            print(f"ERROR segment {segment_idx}: {e}", file=sys.stderr)
            traceback.print_exc()

            error_row = {
                "segment_idx": segment_idx,
                "segment_path": str(segment.segment_path),
                "status": "error",
                "error": repr(e),
            }

            rows = [r for r in rows if int(r["segment_idx"]) != segment_idx]
            rows.append(error_row)
            save_checkpoint(rows, predictions_csv)

            # Continue to the next segment rather than losing the whole run.
            continue

    save_checkpoint(rows, predictions_csv)

    ok_count = sum(1 for r in rows if r.get("status") == "ok")
    err_count = sum(1 for r in rows if r.get("status") == "error")

    print("\nFinished")
    print(f"Successful predictions: {ok_count}")
    print(f"Errors: {err_count}")
    print(f"Saved to: {predictions_csv}")


def parse_args():
    parser = argparse.ArgumentParser(
        description="Run Qwen2.5-VL zero-shot 3RS prediction over video segments."
    )
    parser.add_argument(
        "--segments-csv",
        required=True,
        help="CSV containing segment_idx and segment_path.",
    )
    parser.add_argument(
        "--features-csv",
        required=True,
        help="CSV containing segment_idx and extracted video features.",
    )
    parser.add_argument(
        "--output-dir",
        default="./qwen_3rs_batch_output",
        help="Folder for predictions and logs.",
    )
    parser.add_argument(
        "--official-manual",
        required=True,
        help="Path to the official 3RS v2022 PDF manual.",
    )
    parser.add_argument(
        "--project-guide",
        required=True,
        help="Path to the project 1-minute DOCX coding guide.",
    )
    parser.add_argument(
        "--max-segments",
        type=int,
        default=None,
        help="Optional limit for a smoke test.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()

    run_batch(
        segments_csv=args.segments_csv,
        features_csv=args.features_csv,
        output_dir=args.output_dir,
        official_manual=args.official_manual,
        project_guide=args.project_guide,
        max_segments=args.max_segments,
    )